"""Client Windows : envoie des SMS depuis une colonne Excel via la passerelle Android."""
import argparse
import os
import re
import subprocess
import sys
import time
import uuid

import openpyxl
import requests

PORT = 8765
BASE = f"http://127.0.0.1:{PORT}"
TOKEN = os.environ.get("SMSGW_TOKEN")
if not TOKEN:
    sys.exit("Définissez la variable d'environnement SMSGW_TOKEN (jeton d'appairage affiché sur le téléphone).")
HEADERS = {"Authorization": f"Bearer {TOKEN}"}


def column_index(value):
    result = 0
    for char in value.strip().upper():
        if not "A" <= char <= "Z":
            raise ValueError("La colonne doit être indiquée par sa lettre (ex. A ou C).")
        result = result * 26 + ord(char) - ord("A") + 1
    if not value.strip():
        raise ValueError("La lettre de colonne est vide.")
    return result - 1


def excel_value(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (int, float)) and re.fullmatch(r"0+", cell.number_format or ""):
        return str(int(value)).zfill(len(cell.number_format))
    return str(value).strip()


def is_phone_like(value):
    cleaned = re.sub(r"[\s().-]", "", value.strip())
    return bool(re.fullmatch(r"\+?\d{8,15}", cleaned))


def select_column(ws, requested):
    rows = list(ws.iter_rows())
    if not rows:
        raise ValueError("La feuille Excel est vide.")
    width = max(len(row) for row in rows)
    candidates = []
    for index in range(width):
        values = [excel_value(row[index]) for row in rows[1:] if index < len(row)]
        values = [value for value in values if value]
        matches = sum(is_phone_like(value) for value in values)
        header = excel_value(rows[0][index]) if index < len(rows[0]) else ""
        header_match = bool(re.search(r"phone|tel|téléphone|numero|numéro|number|mobile", header, re.I))
        if matches:
            candidates.append((index, header or f"Colonne {index + 1}", matches / max(1, len(values)) + 0.2 * header_match))
    candidates.sort(key=lambda candidate: candidate[2], reverse=True)

    if requested:
        selected = column_index(requested)
    elif candidates and (len(candidates) == 1 or candidates[0][2] > candidates[1][2]):
        selected = candidates[0][0]
        print(f"Colonne détectée : {openpyxl.utils.get_column_letter(selected + 1)} — {candidates[0][1]}")
    elif candidates:
        print("Plusieurs colonnes contiennent des valeurs ressemblant à des numéros :")
        for index, header, _score in candidates:
            print(f"  {openpyxl.utils.get_column_letter(index + 1)} — {header}")
        selected = column_index(input("Saisissez la lettre de la colonne à utiliser : "))
    else:
        selected = column_index(input("Aucune colonne détectée. Saisissez la lettre de la colonne des numéros : "))

    if selected >= width:
        raise ValueError(f"La colonne {openpyxl.utils.get_column_letter(selected + 1)} n'existe pas dans la feuille.")
    return rows, selected


def api(method, path, **kw):
    response = requests.request(method, BASE + path, headers=HEADERS, timeout=30, **kw)
    try:
        body = response.json()
    except requests.exceptions.JSONDecodeError as exc:
        raise RuntimeError(f"Réponse invalide de la passerelle (HTTP {response.status_code}).") from exc
    if response.status_code >= 400:
        error = body.get("error", "Erreur inconnue") if isinstance(body, dict) else str(body)
        message = error.get("message", str(error)) if isinstance(error, dict) else error
        raise RuntimeError(f"Passerelle HTTP {response.status_code} : {message}")
    return response.status_code, body


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", help="Chemin du classeur Excel .xlsx/.xlsm")
    parser.add_argument("message", help="Texte du SMS à envoyer")
    parser.add_argument("--column", help="Lettre de la colonne des numéros (ex. A, C ou AB)")
    args = parser.parse_args()

    try:
        subprocess.run(["adb", "forward", f"tcp:{PORT}", f"tcp:{PORT}"], check=True)

        workbook = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
        try:
            rows, index = select_column(workbook.active, args.column)
            recipients = [excel_value(row[index]) for row in rows[1:] if index < len(row) and excel_value(row[index])]
        finally:
            workbook.close()
        if not recipients:
            raise ValueError("Aucun numéro trouvé dans la colonne choisie.")

        _, sims = api("GET", "/api/sims")
        selected = next((sim for sim in sims["sims"] if sim["selected"]), None)
        if not selected:
            raise ValueError("Aucune SIM sélectionnée sur le téléphone.")
        print(f"SIM : {selected['displayName']} — {selected['carrierName']} (subscriptionId={selected['subscriptionId']})")

        body = {
            "subscriptionId": selected["subscriptionId"],
            "message": args.message,
            "recipients": recipients,
            "requestId": str(uuid.uuid4()),
        }

        _, report = api("POST", "/api/sms/validate", json=body)
        print(f"{report['rows']} lignes, {report['recipientCount']} uniques, {report['duplicates']} doublons, {report['invalidCount']} invalides")
        print(f"Segments/SMS : {report['segmentsPerSms']} ({report['encoding']}) — nécessaires : {report['smsSegmentsRequired']}, "
              f"restants : {report['smsRemaining']}/{report['smsLimit']}")
        if not report["valid"]:
            for error in report["errors"]:
                print(" ✗", error["code"], "-", error["message"])
            return 1
        if input("Envoyer ? [o/N] ").lower() != "o":
            return 0

        status, job = api("POST", "/api/sms/send", json=body)
        if status not in (200, 202):
            raise RuntimeError(f"Envoi refusé : {job}")
        job_id = job["jobId"]
        print("Job", job_id)
        while True:
            _, progress = api("GET", f"/api/jobs/{job_id}")
            print(f"{progress['status']:<22} envoyés={progress['sent']} échecs={progress['failed']} attente={progress['pending']} ({progress['percent']}%)")
            if progress["status"] in ("COMPLETED", "COMPLETED_WITH_ERRORS", "CANCELLED"):
                break
            time.sleep(3)
        _, failed = api("GET", f"/api/jobs/{job_id}/recipients", params={"status": "FAILED"})
        for recipient in failed["recipients"]:
            print("ÉCHEC", recipient["phoneNumber"], recipient.get("error"))
        return 0
    except (OSError, ValueError, RuntimeError, subprocess.SubprocessError, requests.RequestException) as error:
        print(f"Erreur : {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
